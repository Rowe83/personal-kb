from __future__ import annotations

import os
from typing import List

import fitz
from langchain_core.documents import Document
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".txt", ".md", ".xlsx"})


def normalize_ext(filename: str) -> str:
    _, ext = os.path.splitext(filename or "")
    return ext.lower()


def parse_document(file_path: str, filename: str) -> List[Document]:
    if os.path.getsize(file_path) == 0:
        raise ValueError("上传的文件为空文件 (0字节)")
    ext = normalize_ext(filename)
    if ext == ".pdf":
        return _parse_pdf(file_path, filename)
    if ext == ".txt":
        return _parse_txt(file_path, filename)
    if ext == ".md":
        return _parse_md(file_path, filename)
    if ext == ".xlsx":
        return _parse_xlsx(file_path, filename)
    raise ValueError("不支持的文件类型，仅支持 PDF / TXT / MD / XLSX")


def _parse_pdf(file_path: str, filename: str) -> List[Document]:
    documents: List[Document] = []
    try:
        doc = fitz.open(file_path)
        if doc.page_count == 0:
            raise ValueError("PDF 文件未包含任何页面")

        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text().strip()
            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"filename": filename, "page": page_num + 1},
                    )
                )
        doc.close()
    except Exception as e:
        raise ValueError(f"解析 PDF 文件时发生错误: {e}")

    if not documents:
        raise ValueError("未能从文件中提取出任何有效文件")

    return documents


def _parse_txt(file_path: str, filename: str) -> List[Document]:
    documents: List[Document] = []
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read().strip()
        if not content:
            raise ValueError("TXT 文件内容为空")
        documents.append(
            Document(
                page_content=content, metadata={"filename": filename, "page": 1}
            )
        )
    except UnicodeDecodeError:
        raise ValueError("TXT 文件不是 UTF-8 编码格式")

    return documents


def _parse_md(file_path: str, filename: str) -> List[Document]:
    documents: List[Document] = []
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read().strip()
        if not content:
            raise ValueError("MD 文件内容为空")
        documents.append(
            Document(
                page_content=content, metadata={"filename": filename, "page": 1}
            )
        )
    except UnicodeDecodeError:
        raise ValueError("MD 文件不是 UTF-8 编码格式")

    return documents


def _parse_xlsx(file_path: str, filename: str) -> List[Document]:
    documents: List[Document] = []
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            for sheet_idx, sheet in enumerate(wb.worksheets, start=1):
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                headers = [str(c).strip() if c is not None else "" for c in rows[0]]
                lines: List[str] = []
                for row in rows[1:]:
                    parts = []
                    for i, cell in enumerate(row):
                        if cell is None or str(cell).strip() == "":
                            continue
                        key = headers[i] if i < len(headers) and headers[i] else f"列{i+1}"
                        parts.append(f"{key}: {cell}")
                    if parts:
                        lines.append("\t".join(parts))
                if not lines:
                    continue
                text = "\n".join(lines).strip()
                if not text:
                    continue
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "filename": filename,
                            "page": sheet_idx,
                            "sheet_name": sheet.title,
                        },
                    )
                )
        finally:
            wb.close()
    except Exception as e:
        raise ValueError(f"解析 Excel 文件时发生错误: {e}") from e
    if not documents:
        raise ValueError("未能从 Excel 中提取出任何有效内容")
    return documents
